import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\DELL\Downloads\Daily_OPT_Analysis_Report_2026-09-01_to_2026-09-28.xlsx', data_only=True)
sheet = wb.active

branches = [sheet.cell(row=2, column=c).value for c in range(3, 23)]
new_tot = [sheet.cell(row=15, column=c).value for c in range(3, 23)]
ivf = [sheet.cell(row=16, column=c).value for c in range(3, 23)]
inf = [sheet.cell(row=17, column=c).value for c in range(3, 23)]
anc = [sheet.cell(row=18, column=c).value for c in range(3, 23)]
gyn = [sheet.cell(row=19, column=c).value for c in range(3, 23)]
doc = [sheet.cell(row=20, column=c).value for c in range(3, 23)]
pt = [sheet.cell(row=21, column=c).value for c in range(3, 23)]
onl = [sheet.cell(row=22, column=c).value for c in range(3, 23)]
led = [sheet.cell(row=23, column=c).value for c in range(3, 23)]
off = [sheet.cell(row=24, column=c).value for c in range(3, 23)]

header = "{:<14} | {:<5} | {:<4} | {:<4} | {:<4} | {:<4} | {:<6} | {:<5} | {:<4} | {:<4} | {:<4}".format(
    "Branch", "Total", "IVF", "INF", "ANC", "GYN", "Online", "Lead", "Doc", "Pt", "Off"
)
print(header)
print("-" * len(header))
for i in range(len(branches)):
    line = "{:<14} | {:<5} | {:<4} | {:<4} | {:<4} | {:<4} | {:<6} | {:<5} | {:<4} | {:<4} | {:<4}".format(
        str(branches[i]), str(new_tot[i]), str(ivf[i]), str(inf[i]), str(anc[i]), str(gyn[i]),
        str(onl[i]), str(led[i]), str(doc[i]), str(pt[i]), str(off[i])
    )
    print(line)
